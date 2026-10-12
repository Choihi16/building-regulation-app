"""
구조화 문서 보고서 생성 모듈 (Document Report Generator)
- python-docx를 활용하여 한글(HWP) 및 MS Word 호환 정규 보고서(.docx) 생성
- 경영진 보고용 3줄 결론, 토지대장 정합성, 세대구성 시뮬레이션, 4대 심의 판정표,
  지구단위계획 6대 핵심 지침 실무 체크리스트 정규 표 포함
- 표 서식, 헤더 셀 음영(Navy), 폰트 크기 계층, 콜아웃 박스 등 서식 적용
"""

import io
from typing import Dict, Any, List
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# 색상 상수
COLOR_NAVY = RGBColor(0x1E, 0x3A, 0x8A)
COLOR_BLUE = RGBColor(0x25, 0x63, 0xEB)
COLOR_DARK = RGBColor(0x0F, 0x17, 0x2A)
COLOR_TEXT = RGBColor(0x1E, 0x29, 0x3B)
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX_NAVY = "1E3A8A"
HEX_BLUE = "2563EB"
HEX_SLATE = "475569"
HEX_LIGHT_BLUE = "EFF6FF"
HEX_LIGHT_GRAY = "F8FAFC"
HEX_LIGHT_YELLOW = "FEF08A"
HEX_LIGHT_GREEN = "DEF7EC"
HEX_LIGHT_RED = "FDE8E8"

def _set_cell_background(cell, hex_color: str):
    """표 셀의 배경색을 지정"""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def _set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """셀 내부 여백(padding) 설정 (dxa 단위: 1pt = 20dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def _format_cell(cell, text: str, font_size=9, bold=False, color=COLOR_TEXT,
                 align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=None, padding_top=100, padding_bottom=100):
    """셀 텍스트 및 서식 일괄 적용 헬퍼"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.font.name = "맑은 고딕"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if bg_color:
        _set_cell_background(cell, bg_color)
    _set_cell_margins(cell, top=padding_top, bottom=padding_bottom, left=140, right=140)

def _style_header_row(row, headers: List[str], bg_color=HEX_NAVY, font_size=9.5):
    """표의 헤더 행 스타일링"""
    for idx, header in enumerate(headers):
        cell = row.cells[idx]
        _format_cell(cell, header, font_size=font_size, bold=True, color=COLOR_WHITE,
                     align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg_color, padding_top=140, padding_bottom=140)

def _add_section_heading(doc, text: str, level=1):
    """문서 섹션 제목 추가"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_NAVY
    elif level == 2:
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_BLUE
    else:
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_DARK

def generate_docx_report(analysis: Dict[str, Any]) -> io.BytesIO:
    """분석 결과 딕셔너리를 기반으로 정규 Word/HWP 호환 서식 보고서 생성"""
    doc = Document()

    # 페이지 여백 설정 (상하 20mm, 좌우 20mm)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    site = analysis.get("site", {})
    land = analysis.get("land", {})
    zoning = analysis.get("zoning_info", {})
    planned = analysis.get("planned_conditions", {})
    area_val = analysis.get("area_validation", {})
    bcr = analysis.get("bcr", {})
    far = analysis.get("far", {})
    feas = analysis.get("feasibility", {})
    height = analysis.get("height", {})
    delibs = analysis.get("deliberations", {})
    constraints = analysis.get("design_constraints", {})
    special = analysis.get("special_ordinance", {})
    dp_alert = analysis.get("district_plan_alert", {})
    three_lines = analysis.get("executive_summary_3lines", [])
    parking_info = constraints.get("parking", {})

    # =========================================================================
    # 표지 및 메인 타이틀
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r_sub = p_title.add_run("[경영진 및 임원진 보고용]\n")
    r_sub.font.name = "맑은 고딕"
    r_sub.font.size = Pt(10)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_BLUE

    r_main = p_title.add_run("🏛️ 건축 법규 및 기획설계 사업성 통합 검토 보고서")
    r_main.font.name = "맑은 고딕"
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = COLOR_NAVY

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(2)
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run(
        f"• 대상 부지: {site.get('refined_address', site.get('raw_address'))}\n"
        f"• 필지고유번호(PNU): {site.get('pnu', '확인필요')} | 법정용도지역: {land.get('primary_zoning', '미지정')}\n"
        f"• 검토 일자: 2026-10-12 | 작성: 개발기획팀 / AI 건축법규 분석에이전트 (공부상 토지대장 검증 완료)"
    )
    r_meta.font.name = "맑은 고딕"
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # 1. 경영진 보고용 3줄 결론 (Callout Box)
    # =========================================================================
    _add_section_heading(doc, "1. 💡 경영진 보고용 3줄 핵심 결론 (Executive Summary)", level=1)
    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.style = "Table Grid"
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = callout_tbl.cell(0, 0)
    _set_cell_background(c_cell, HEX_LIGHT_BLUE)
    _set_cell_margins(c_cell, top=160, bottom=160, left=200, right=200)
    c_p = c_cell.paragraphs[0]
    c_p.paragraph_format.space_before = Pt(0)
    c_p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(three_lines, 1):
        if idx > 1:
            c_p = c_cell.add_paragraph()
            c_p.paragraph_format.space_before = Pt(4)
            c_p.paragraph_format.space_after = Pt(0)
        r = c_p.add_run(f"[{idx}] {line}")
        r.font.name = "맑은 고딕"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

    # =========================================================================
    # 2. 핵심 기획설계 및 사업성 지표 요약표
    # =========================================================================
    _add_section_heading(doc, "2. 📊 핵심 기획설계 및 사업성 지표 (Executive KPIs)", level=1)
    kpi_tbl = doc.add_table(rows=1, cols=4)
    kpi_tbl.style = "Table Grid"
    kpi_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_kpi = ["항목 구분", "적용 기준치", "산출 면적 / 규제 수치", "평형 환산 / 비고"]
    _style_header_row(kpi_tbl.rows[0], headers_kpi, bg_color=HEX_NAVY)

    kpi_data = [
        ("① 적용 건폐율 (BCR)", bcr.get("applied_limit"), f"{feas.get('max_building_area_sqm',0):,.1f} ㎡", f"{bcr.get('comparison_summary')} | 최대 건축면적 약 {feas.get('max_building_area_py',0):,.1f} 평"),
        ("② 적용 용적률 (FAR)", far.get("applied_limit"), f"{feas.get('max_ground_gfa_sqm',0):,.1f} ㎡", f"{far.get('comparison_summary')} | 지상 연면적 약 {feas.get('max_ground_gfa_py',0):,.1f} 평 (인센티브 시 {feas.get('incentive_far_pct')}%)"),
        ("③ 예상 최고높이 / 층수", height.get("summary_limit"), f"지상 {feas.get('est_floors')} 층", f"높이 약 {feas.get('est_height_m')} m (층당 {feas.get('floor_height_m')}m 기준)"),
        ("④ 주요 심의 리스크", delibs.get("risk_summary", {}).get("overall_grade"), f"총 {delibs.get('risk_summary', {}).get('high_risk_count')}개 대상", f"예상 사전인허가 기간: {delibs.get('risk_summary', {}).get('total_timeline_est')}"),
        ("⑤ 예상 총 연면적", "지상 100% + 지하 45%", f"{feas.get('est_total_gfa_sqm',0):,.1f} ㎡", f"총 약 {feas.get('est_total_gfa_py',0):,.1f} 평 (지하 약 {feas.get('est_basement_gfa_sqm',0):,.1f}㎡)"),
        ("⑥ 계획 세대 / 주차", f"{planned.get('unit_count')} 세대/호", f"법정 {parking_info.get('legal_count')}대 / 권장 {parking_info.get('recommended_count')}대", f"친환경 충전 {parking_info.get('ev_charging_count')}대 포함")
    ]
    for row_idx, kd in enumerate(kpi_data):
        row = kpi_tbl.add_row()
        bg = HEX_LIGHT_GRAY if row_idx % 2 == 1 else "FFFFFF"
        _format_cell(row.cells[0], kd[0], font_size=9, bold=True, color=COLOR_DARK, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=bg)
        _format_cell(row.cells[1], kd[1], font_size=9, bold=False, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg)
        _format_cell(row.cells[2], kd[2], font_size=9, bold=True, color=COLOR_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg)
        _format_cell(row.cells[3], kd[3], font_size=8.5, bold=False, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=bg)

    # =========================================================================
    # 3. 대지 지적 속성 및 토지대장 공부상 면적 검증
    # =========================================================================
    _add_section_heading(doc, "3. 📍 대지 지적 속성 및 토지대장 공부상 면적 검증", level=1)
    prop_tbl = doc.add_table(rows=1, cols=4)
    prop_tbl.style = "Table Grid"
    prop_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _style_header_row(prop_tbl.rows[0], ["구분 항목", "세부 내역", "구분 항목", "세부 내역"], bg_color=HEX_SLATE)

    match_str = "일치 (정상)" if area_val.get("is_match") else f"⚠️ 편차 {area_val.get('diff_sqm'):,.1f}㎡ ({area_val.get('diff_pct')}%)"
    prop_data = [
        ("소재지 주소", site.get("refined_address", ""), "필지고유번호(PNU)", site.get("pnu", "")),
        ("법정 용도지역", land.get("primary_zoning", ""), "행정구역", f"{site.get('sido')} {site.get('sigg')} {site.get('emd')}"),
        ("검토 대지면적", f"{planned.get('site_area_sqm',0):,.1f} ㎡ (약 {planned.get('site_area_py',0):,.1f}평)", "토지대장 공부면적", f"{area_val.get('official_area_sqm',0):,.1f} ㎡ (약 {area_val.get('official_area_py',0):,.1f}평)"),
        ("면적 정합성 판정", match_str, "공부상 지목/지가", f"{area_val.get('jimok')} / {area_val.get('official_price')}"),
        ("도로접면 / 지세", f"{area_val.get('road_side')} / {area_val.get('topography_height')}", "지구단위계획 여부", f"포함 ({', '.join(land.get('district_plan_names', []))})" if land.get("is_district_plan") else "해당없음")
    ]
    for p_row in prop_data:
        r = prop_tbl.add_row()
        _format_cell(r.cells[0], p_row[0], font_size=8.5, bold=True, color=COLOR_DARK, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=HEX_LIGHT_GRAY)
        _format_cell(r.cells[1], p_row[1], font_size=8.5, bold=False, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.LEFT)
        _format_cell(r.cells[2], p_row[2], font_size=8.5, bold=True, color=COLOR_DARK, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=HEX_LIGHT_GRAY)
        _format_cell(r.cells[3], p_row[3], font_size=8.5, bold=False, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.LEFT)

    # =========================================================================
    # 4. 법정 vs 지자체 조례 제한 비교표
    # =========================================================================
    _add_section_heading(doc, "4. ⚖️ 법정 vs 지자체 조례/지구단위계획 제한 이중 비교표", level=1)
    comp_tbl = doc.add_table(rows=1, cols=5)
    comp_tbl.style = "Table Grid"
    comp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _style_header_row(comp_tbl.rows[0], ["규제 항목", "최종 적용 (엄격)", "지자체 조례/지침 상한", "국토계획법 상한", "이중 비교 요약 및 근거"], bg_color=HEX_NAVY)

    comp_data = [
        ("건폐율 (BCR)", f"★ {bcr.get('applied_limit')}", bcr.get("ordinance_limit"), bcr.get("national_limit"), f"{bcr.get('comparison_summary')} / {bcr.get('law_source')}"),
        ("용적률 (FAR)", f"★ {far.get('applied_limit')}", far.get("ordinance_limit"), far.get("national_limit"), f"{far.get('comparison_summary')} / {far.get('law_source')}"),
        ("최고높이", f"지상 {feas.get('est_floors')}층", height.get("summary_limit"), "건축법 제60/61조", f"가로구역 및 일조사선 동시 충족 / {site.get('sido')} 건축조례")
    ]
    for cd in comp_data:
        r = comp_tbl.add_row()
        _format_cell(r.cells[0], cd[0], font_size=8.5, bold=True, color=COLOR_DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[1], cd[1], font_size=9, bold=True, color=COLOR_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[2], cd[2], font_size=8.5, bold=False, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[3], cd[3], font_size=8.5, bold=False, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[4], cd[4], font_size=8, bold=False, color=COLOR_MUTED, align=WD_ALIGN_PARAGRAPH.LEFT)

    # =========================================================================
    # 5. 용도별 설계 제약, 세대 구성 및 주차 산정
    # =========================================================================
    _add_section_heading(doc, f"5. 🏢 [{planned.get('planned_usage')}] 설계 제약, 세대구성 및 주차 산정", level=1)

    p_rule = doc.add_paragraph()
    p_rule.paragraph_format.space_before = Pt(4)
    p_rule.paragraph_format.space_after = Pt(4)
    p_rule.add_run("• 계획 용도 필수 규정: ").bold = True
    for r_idx, rule in enumerate(constraints.get("usage_rules", [])):
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(2)
        r_item = p_sub.add_run(f"  - {rule}")
        r_item.font.name = "맑은 고딕"
        r_item.font.size = Pt(8.5)

    # 세대 구성표
    breakdown = feas.get("unit_breakdown", {})
    if breakdown:
        _add_section_heading(doc, "5.1 세대 및 호실 세부 구성 시뮬레이션", level=2)
        unit_tbl = doc.add_table(rows=1, cols=5)
        unit_tbl.style = "Table Grid"
        unit_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        _style_header_row(unit_tbl.rows[0], ["용도 구분", "평형 타입", "전용면적(㎡)", "평형 환산", "계획 세대/호수"], bg_color=HEX_SLATE)

        if breakdown.get("apt"):
            for k, v in breakdown["apt"].items():
                if v > 0:
                    r = unit_tbl.add_row()
                    _format_cell(r.cells[0], "공동주택 (아파트)", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[1], f"{k} 타입", font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[2], f"{float(k):,.1f} ㎡", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[3], f"약 {round(float(k)/3.305785, 1)} 평", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[4], f"{v} 세대", font_size=8.5, bold=True, color=COLOR_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)

        if breakdown.get("officetel"):
            for k, v in breakdown["officetel"].items():
                if v > 0:
                    r = unit_tbl.add_row()
                    _format_cell(r.cells[0], "오피스텔 (업무시설)", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[1], f"{k} 타입", font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[2], f"{float(k):,.1f} ㎡", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[3], f"약 {round(float(k)/3.305785, 1)} 평", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                    _format_cell(r.cells[4], f"{v} 호", font_size=8.5, bold=True, color=COLOR_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)

        r_tot = unit_tbl.add_row()
        _format_cell(r_tot.cells[0], "합계", font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=HEX_LIGHT_BLUE)
        _format_cell(r_tot.cells[1], "-", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=HEX_LIGHT_BLUE)
        _format_cell(r_tot.cells[2], "-", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=HEX_LIGHT_BLUE)
        _format_cell(r_tot.cells[3], "-", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=HEX_LIGHT_BLUE)
        _format_cell(r_tot.cells[4], f"총 {feas.get('effective_unit_count')} 세대/호", font_size=8.5, bold=True, color=COLOR_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=HEX_LIGHT_BLUE)

    # 주차 산정표
    _add_section_heading(doc, "5.2 법정 및 조례 기준 주차대수 산정 결과", level=2)
    p_tbl = doc.add_table(rows=1, cols=4)
    p_tbl.style = "Table Grid"
    p_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _style_header_row(p_tbl.rows[0], ["구분 항목", "산정 대수", "적용 근거 법령", "비고 및 필수 충족 요건"], bg_color=HEX_BLUE)

    park_data = [
        ("법정 최소 주차대수", f"{parking_info.get('legal_count')} 대", "주차장법 시행령 별표1", "시설면적 및 전용면적 기준 법정 최소치"),
        ("지자체 조례 기준", f"{parking_info.get('ordinance_count')} 대", f"{site.get('sido')} 주차장 조례", "지자체 강화 기준 적용"),
        ("★ 실무 권장 주차대수", f"{parking_info.get('recommended_count')} 대", "건축심의 권고 / 분양성", "세대당 1.1~1.3대 슬라이딩 적용치"),
        ("친환경차 충전시설", f"{parking_info.get('ev_charging_count')} 대", "친환경차법 제11조의2", "총 주차대수의 5% 이상 의무화"),
        ("장애인 전용 주차", f"{parking_info.get('accessibility_count')} 대", "장애인등편의법", "법정 2~4% 범위 의무 구획")
    ]
    for pd in park_data:
        r = p_tbl.add_row()
        _format_cell(r.cells[0], pd[0], font_size=8.5, bold=True, color=COLOR_DARK, align=WD_ALIGN_PARAGRAPH.LEFT)
        _format_cell(r.cells[1], pd[1], font_size=9, bold=True, color=COLOR_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[2], pd[2], font_size=8.5, bold=False, color=COLOR_MUTED, align=WD_ALIGN_PARAGRAPH.LEFT)
        _format_cell(r.cells[3], pd[3], font_size=8.5, bold=False, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.LEFT)

    # =========================================================================
    # 6. 4대 주요 심의 자동 판정 및 소방 안전
    # =========================================================================
    _add_section_heading(doc, "6. 🛡️ 4대 주요 심의 자동 판정 및 소방 안전 매뉴얼", level=1)
    delib_tbl = doc.add_table(rows=1, cols=6)
    delib_tbl.style = "Table Grid"
    delib_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _style_header_row(delib_tbl.rows[0], ["심의 명칭", "대상여부", "등급", "소요기간", "주관기관", "핵심 체크포인트"], bg_color=HEX_NAVY)

    for d in delibs.get("items", []):
        r = delib_tbl.add_row()
        _format_cell(r.cells[0], d.get("name"), font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        is_req = "대상" in str(d.get("required"))
        bg_req = HEX_LIGHT_RED if d.get("tag") == "규제" else (HEX_LIGHT_YELLOW if d.get("tag") == "주의" else HEX_LIGHT_GREEN)
        _format_cell(r.cells[1], d.get("required"), font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg_req)
        _format_cell(r.cells[2], d.get("tag"), font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg_req)
        _format_cell(r.cells[3], d.get("timeline"), font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[4], d.get("authority"), font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[5], d.get("checkpoints"), font_size=8, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 소방 안전 지침
    p_fire = doc.add_paragraph()
    p_fire.paragraph_format.space_before = Pt(6)
    p_fire.paragraph_format.space_after = Pt(2)
    p_fire.add_run("• 피난·방화·소방안전 필수 착안사항:").bold = True
    for fn in delibs.get("fire_safety_notes", []):
        pf = doc.add_paragraph()
        pf.paragraph_format.space_before = Pt(0)
        pf.paragraph_format.space_after = Pt(2)
        r_f = pf.add_run(f"  - {fn}")
        r_f.font.name = "맑은 고딕"
        r_f.font.size = Pt(8.5)

    # =========================================================================
    # 7. 지구단위계획 6대 핵심 지침 실무 체크리스트 (정규 표)
    # =========================================================================
    _add_section_heading(doc, "7. 📋 지구단위계획 6대 핵심 지침 실무 점검 체크리스트", level=1)

    p_dp_notice = doc.add_paragraph()
    p_dp_notice.paragraph_format.space_before = Pt(2)
    p_dp_notice.paragraph_format.space_after = Pt(4)
    r_dpn = p_dp_notice.add_run("※ 유의사항: 지구단위계획구역은 일반 건축조례보다 지구단위계획 수립 지침이 최우선 적용됩니다. 아래 6대 기준을 사전 확인하여 설계에 반영해야 합니다.")
    r_dpn.font.name = "맑은 고딕"
    r_dpn.font.size = Pt(8.5)
    r_dpn.font.color.rgb = COLOR_BLUE

    dp_tbl = doc.add_table(rows=1, cols=5)
    dp_tbl.style = "Table Grid"
    dp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    dp_headers = ["구분 카테고리", "세부 규정 항목", "법적 및 지침 기준", "실무 설계 점검사항", "인센티브 및 완화 조건"]
    _style_header_row(dp_tbl.rows[0], dp_headers, bg_color=HEX_NAVY)

    for chk in special.get("district_plan_checklist", []):
        r = dp_tbl.add_row()
        _format_cell(r.cells[0], chk.get("카테고리"), font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=HEX_LIGHT_GRAY)
        _format_cell(r.cells[1], chk.get("세부 규정 항목"), font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _format_cell(r.cells[2], chk.get("법적 및 지침 기준"), font_size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
        _format_cell(r.cells[3], chk.get("실무 설계 점검사항"), font_size=8, bold=True, color=COLOR_BLUE, align=WD_ALIGN_PARAGRAPH.LEFT)
        _format_cell(r.cells[4], chk.get("완화 및 인센티브"), font_size=8, align=WD_ALIGN_PARAGRAPH.LEFT)

    # =========================================================================
    # 8. 착수 전 필수 리스크 체크리스트
    # =========================================================================
    _add_section_heading(doc, "8. 🚨 사업 착수 전 필수 리스크 체크리스트", level=1)
    risk_tbl = doc.add_table(rows=1, cols=2)
    risk_tbl.style = "Table Grid"
    risk_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _style_header_row(risk_tbl.rows[0], ["체크 항목", "착수 전 선행 실무 조치 사항"], bg_color=HEX_SLATE)

    risks = [
        ("일조권 사선제한 검토", "정북방향 인접대지의 레벨(고저차) 및 경계선 확인을 위한 현황측량 선행 필수"),
        ("지하 굴토심의 리스크", "지하 2층 이상 굴착 시 안전관리계획서 및 지반조사보고서(시추조사) 조기 착수 필수"),
        ("지구단위계획 지침 준수", "건축선 후퇴(전면공지), 공공보행통로 폭(4~6m) 및 1층 가로활성화 투시벽면율 50% 사전 반영"),
        ("교통영향평가 대상성", "연면적 변경에 따른 교통심의 트리거 규모(2만~3.6만㎡) 도달 여부 모니터링"),
        ("전기차 화재방화 기준", "지하주차장 스프링클러 헤드 간격(2.3m) 및 방화구획 최신 소방지침 반영")
    ]
    for rk in risks:
        r = risk_tbl.add_row()
        _format_cell(r.cells[0], rk[0], font_size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        _format_cell(r.cells[1], rk[1], font_size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 바이트 스트림 변환
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
